import hashlib
import json
from pathlib import Path

import fitz
from openpyxl import Workbook

from tools.bp_v6.content import SLIDES
from tools.bp_v6.theme import add_text, new_presentation
from tools.bp_v6.verify_artifacts import verify_artifacts


PDF_NAME = "逸飞AI智眼系统_创业大赛BP_V6.pdf"


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest().upper()


def _pptx(path: Path) -> None:
    prs = new_presentation()
    for spec in SLIDES:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(slide, spec.title, 0.5, 0.5, 12, 0.5)
    add_text(prs.slides[3], "公司成立以来收入\n0万元", 0.5, 1.5, 3, 1)
    add_text(prs.slides[8], "P95≤10秒\n试点验收目标，不是历史业绩", 0.5, 1.5, 5, 1)
    add_text(prs.slides[14], "融资需求\n500万元", 0.5, 1.5, 3, 1)
    prs.save(path)


def _pdf(path: Path) -> None:
    document = fitz.open()
    for index in range(16):
        page = document.new_page(width=960, height=540)
        page.insert_text((80, 100), f"BP V6 PAGE {index + 1} P95", fontsize=24)
    document.save(path)
    document.close()


def _workbook(path: Path, link: bool) -> None:
    workbook = Workbook()
    workbook.active.title = "项目模板"
    for index in range(2, 12):
        workbook.create_sheet(f"Sheet{index}")
    if link:
        cell = workbook.worksheets[0]["J3"]
        cell.value = PDF_NAME
        cell.hyperlink = PDF_NAME
    workbook.save(path)


def test_verifies_complete_artifact_contract(tmp_path):
    source_pdf = tmp_path / "source.pdf"
    source_pdf.write_bytes(b"immutable pdf source")
    source_workbook = tmp_path / "source.xlsx"
    _workbook(source_workbook, link=False)
    pptx = tmp_path / "deck.pptx"
    _pptx(pptx)
    pdf = tmp_path / PDF_NAME
    _pdf(pdf)
    final_workbook = tmp_path / "final.xlsx"
    _workbook(final_workbook, link=True)
    baseline = tmp_path / "baseline.json"
    baseline.write_text(
        json.dumps(
            {
                "source_pdf": {"path": str(source_pdf), "sha256": _hash(source_pdf), "pages": 16},
                "source_workbook": {
                    "path": str(source_workbook),
                    "sha256": _hash(source_workbook),
                    "sheet_count": 11,
                    "first_sheet": "项目模板",
                    "pdf_cell": "J3",
                },
            }
        ),
        encoding="utf-8",
    )

    report = verify_artifacts(pptx, pdf, final_workbook, baseline)

    assert report["status"] == "passed"
    assert report["pptx"]["slides"] == 16
    assert report["pdf"]["pages"] == 16
    assert report["workbook"]["pdf_link"] == PDF_NAME
    assert report["sources_unchanged"] is True
